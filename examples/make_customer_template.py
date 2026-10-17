"""専用様式の取り込みを試すための、架空の顧客様式（テーブル定義書）を作る。

  python3 examples/make_customer_template.py examples/顧客様式_テーブル定義書.xlsx
実在の顧客の様式ではない。よくある形（表紙・改訂履歴・テーブルごとのシート）を真似たもの。
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Border, Font, PatternFill, Side, Alignment

THIN = Side(style="thin")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD = PatternFill("solid", fgColor="D9D9D9")


def make(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "表紙"
    ws["B2"], ws["B2"].font = "テーブル定義書", Font(size=20, bold=True)
    for r, label in enumerate(["システム名", "作成者", "作成日", "版", "承認者"], 5):
        ws[f"B{r}"] = label
        ws[f"B{r}"].font, ws[f"B{r}"].fill, ws[f"B{r}"].border = Font(bold=True), HEAD, BOX
        ws.merge_cells(f"C{r}:E{r}")
        for c in "CDE":
            ws[f"{c}{r}"].border = BOX
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 20
    h = wb.create_sheet("改訂履歴")
    for c, label in zip("ABCD", ["版", "日付", "変更内容", "変更者"]):
        h[f"{c}3"], h[f"{c}3"].fill, h[f"{c}3"].border, h[f"{c}3"].font = label, HEAD, BOX, Font(bold=True)
    for r in range(4, 21):
        for c in "ABCD":
            h[f"{c}{r}"].border = BOX
    h.column_dimensions["C"].width = 60
    t = wb.create_sheet("テーブル定義")
    for r, label in enumerate(["テーブル論理名", "テーブル物理名", "説明"], 2):
        t[f"B{r}"], t[f"B{r}"].fill, t[f"B{r}"].border, t[f"B{r}"].font = label, HEAD, BOX, Font(bold=True)
        t.merge_cells(f"C{r}:F{r}")
        for c in "CDEF":
            t[f"{c}{r}"].border = BOX
    heads = ["No", "論理名", "物理名", "データ型", "桁数", "NOT NULL", "PK", "デフォルト", "区分", "備考"]
    for n, label in enumerate(heads):
        c = chr(ord("A") + n)
        t[f"{c}6"], t[f"{c}6"].fill, t[f"{c}6"].border, t[f"{c}6"].font = label, HEAD, BOX, Font(bold=True)
        t[f"{c}6"].alignment = Alignment(horizontal="center")
        for r in range(7, 37):
            t[f"{c}{r}"].border = BOX
    for c, w in zip("ABCDEFGHIJ", [5, 18, 18, 12, 6, 9, 5, 12, 10, 30]):
        t.column_dimensions[c].width = w
    wb.save(path)


if __name__ == "__main__":
    make(sys.argv[1] if len(sys.argv) > 1 else "顧客様式_テーブル定義書.xlsx")
