"""台帳からExcelを清書する（banto export）。

台帳MDが原本で、Excelは提出用の写し。本体のシートは保護（パスワードなし）して直接の編集を防ぎ、
指摘はレビュー表に書いてもらう。各行には台帳への対応（banto-ref）を非表示列で埋め込む。
"""
import datetime
import hashlib
import json
import os
import re
import subprocess
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from . import engine, ledger as LG
from .domains import get
from . import fio

FONT = "Yu Gothic"
HEAD_FILL = PatternFill("solid", fgColor="DCE6F1")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
OPEN_FILL = PatternFill("solid", fgColor="FCE4D6")
THIN = Side(style="thin", color="A6A6A6")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
SEV_JA = {"high": "高", "medium": "中", "low": "低"}
STATUS_JA = {"open": "未判断", "decided": "直し待ち", "resolved": "解決済み"}
CONTAINER_STATUS_JA = {"draft": "未確定あり", "confirmed": "確定", "pending": "有無が未判断", "removing": "削除予定",
                       "removed": "削除済み"}


def _f(bold=False, size=10, color=None):
    return Font(name=FONT, bold=bold, size=size, color=color)


def _sheet_title(name, used):
    t = re.sub(r"[\[\]:*?/\\]", "_", name)[:31] or "sheet"
    base, n = t, 2
    while t in used:
        suffix = f"~{n}"
        t = base[:31 - len(suffix)] + suffix
        n += 1
    used.add(t)
    return t


def _table(ws, row, headers, rows, widths=None, fills=None):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font, cell.fill, cell.border = _f(True), HEAD_FILL, BOX
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for r, vals in enumerate(rows, row + 1):
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font, cell.border = _f(), BOX
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if fills and fills[r - row - 1]:
                cell.fill = fills[r - row - 1]
    if widths:
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = w
    return row + len(rows) + 1


def _print_setup(ws, title):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddHeader.center.text = title
    ws.oddFooter.center.text = "&P / &N"


def _git_head(project):
    try:
        r = subprocess.run(["git", "-C", project, "rev-parse", "--short", "HEAD"], capture_output=True, text=True, timeout=30)
        return r.stdout.strip() or "（Git管理外）"
    except Exception:
        return "（Git管理外）"


def _versions(res, dname):
    """台帳の中身が前回の出力から変わっていれば版を上げ、改訂履歴に残す。"""
    p = os.path.join(res.base, "state", "versions.json")
    data = json.loads(fio.read(p)) if os.path.exists(p) else {}
    dom = get(dname)
    d = os.path.join(res.base, dom.ledger_dir)
    h = hashlib.sha256()
    for f in sorted(os.listdir(d)) if os.path.isdir(d) else []:
        h.update(fio.read_bytes(os.path.join(d, f)))
    digest = h.hexdigest()[:16]
    v = data.setdefault(dname, {"version": 0, "hash": "", "history": [], "decisions": []})
    if v["hash"] != digest:
        applied = sorted(set(res.state.applied) - set(v["decisions"]))
        v["version"] += 1
        v["hash"] = digest
        v["history"].append({"version": v["version"], "date": datetime.date.today().isoformat(),
                             "decisions": applied,
                             "note": "初版" if v["version"] == 1 else ("判断の反映：" + "、".join(applied) if applied else "台帳の更新")})
        v["decisions"] = sorted(res.state.applied)
        os.makedirs(os.path.dirname(p), exist_ok=True)
        fio.write(p, json.dumps(data, ensure_ascii=False, indent=1, sort_keys=True) + "\n")
    return v


def export(project, out=None):
    res = engine.sync(project)
    out = out or os.path.join(res.base, "export")
    os.makedirs(out, exist_ok=True)
    head = _git_head(project)
    files, versions = [], {}
    for dname, ledger in res.ledgers.items():
        dom = get(dname)
        v = _versions(res, dname)
        versions[dom.title] = v["version"]
        path = os.path.join(out, f"{dname}.xlsx")
        _domain_book(res, dom, ledger, v, head, path)
        files.append(path)
    files.append(_review_book(res, versions, out))
    return files


def _domain_book(res, dom, ledger, ver, head, path):
    wb = Workbook()
    used = set()
    drifts = [d for d in res.drifts if d.domain is dom]
    open_n = sum(1 for d in drifts if d.status == "open")
    dec_n = sum(1 for d in drifts if d.status == "decided")
    # 表紙
    ws = wb.active
    ws.title = _sheet_title("表紙", used)
    ws["B3"], ws["B3"].font = dom.title, _f(True, 20)
    info = [("プロジェクト", os.path.basename(os.path.abspath(res.project))), ("版", f"第{ver['version']}版"),
            ("作成日", datetime.date.today().isoformat()), ("生成元", f"banto台帳（{dom.ledger_dir}）／コミット {head}"),
            (f"{dom.container_label}の数", len([c for c in ledger.values() if c.status != "removed"])),
            (f"{dom.item_label}の数", sum(len(c.items) for c in ledger.values())),
            ("未判断のずれ", open_n), ("直し待ちのずれ", dec_n)]
    for n, (k, val) in enumerate(info, 6):
        ws.cell(row=n, column=2, value=k).font = _f(True)
        c = ws.cell(row=n, column=3, value=val)
        c.font, c.alignment = _f(), Alignment(horizontal="left")
        if k in ("未判断のずれ", "直し待ちのずれ") and val:
            c.fill = WARN_FILL
    for n, text in enumerate(["このファイルはbantoの台帳（MD）から自動生成した写しです。直接編集しないでください。",
                              "指摘は、同時に出力したレビュー表に、各シートのIDを添えて記入してください。"], 16):
        ws.cell(row=n, column=2, value=text).font = _f(color="C00000")
    ws.column_dimensions["B"].width, ws.column_dimensions["C"].width = 18, 60
    _print_setup(ws, dom.title)
    # 改訂履歴
    ws = wb.create_sheet(_sheet_title("改訂履歴", used))
    _table(ws, 1, ["版", "日付", "内容", "判断"],
           [[f"第{h['version']}版", h["date"], h["note"], "、".join(h["decisions"])] for h in ver["history"]],
           [10, 14, 50, 40])
    # 一覧
    ws = wb.create_sheet(_sheet_title("一覧", used))
    rows, fills = [], []
    for n, (name, lc) in enumerate(sorted(ledger.items(), key=lambda kv: kv[1].id), 1):
        no = sum(1 for d in drifts if d.container == name and d.status == "open")
        rows.append([n, lc.id, name, len(lc.items), CONTAINER_STATUS_JA.get(lc.status, lc.status), no])
        fills.append(OPEN_FILL if no else None)
    _table(ws, 1, ["No", "ID", dom.container_label, f"{dom.item_label}の数", "状態", "未判断"], rows, [6, 10, 36, 10, 14, 8], fills)
    ws.freeze_panes = "A2"
    _print_setup(ws, f"{dom.title} 一覧")
    # 本体
    by_key = {d.key: d for d in drifts}
    for name, lc in sorted(ledger.items(), key=lambda kv: kv[1].id):
        if lc.status == "removed":
            continue
        ws = wb.create_sheet(_sheet_title(name, used))
        ws["A1"], ws["A1"].font = f"{dom.container_label}：{name}", _f(True, 14)
        ws["A2"] = f"ID：{lc.id}　状態：{CONTAINER_STATUS_JA.get(lc.status, lc.status)}"
        ws["A2"].font = _f()
        r = 4
        if dom.container_attrs():
            crow = []
            for a in dom.container_attrs():
                v = lc.attrs.get(a.name)
                d = by_key.get(f"{dom.name}|{name}||{a.name}")
                crow.append([a.label, _cell(a, v, d)])
            r = _table(ws, r, ["属性", "値"], crow) + 1
        iattrs = dom.item_attrs()
        heads = ["ID", dom.item_label] + [a.label for a in iattrs] + list(dom.free_columns) + ["状態", "banto-ref"]
        rows, fills = [], []
        lrel = os.path.relpath(os.path.join(res.base, dom.ledger_dir, LG.file_name(name)), res.project).replace(os.sep, "/")
        for iname, it in lc.items.items():
            row = [it.id, iname]
            has_open = False
            for a in iattrs:
                d = by_key.get(f"{dom.name}|{name}|{iname}|{a.name}")
                row.append(_cell(a, it.values.get(a.name), d))
                has_open |= bool(d and d.status == "open")
            row += [it.free.get(fc, "") for fc in dom.free_columns]
            ed = by_key.get(f"{dom.name}|{name}|{iname}|exists")
            st = "有無が未判断" if it.exists is None else ("削除予定" if it.exists is False else "")
            if ed and ed.status != "resolved":
                st += f"（{ed.id}）"
                has_open = True
            row += [st or ("未確定あり" if has_open else "確定"), f"{lrel}#{it.id}"]
            rows.append(row)
            fills.append(OPEN_FILL if has_open or it.exists is None else None)
        widths = [10, 30] + [48 if a.name == "purpose" else max(15, len(a.label) * 2 + 4) for a in iattrs] + \
                 [24 for _ in dom.free_columns] + [16, 30]
        _table(ws, r, heads, rows, widths, fills)
        ws.column_dimensions[get_column_letter(len(heads))].hidden = True
        ws.freeze_panes = ws.cell(row=r + 1, column=3)
        _print_setup(ws, f"{dom.title} {name}")
        ws.protection.sheet = True
    # 許容された差
    ws = wb.create_sheet(_sheet_title("許容された差", used))
    rows = []
    for name, lc in sorted(ledger.items(), key=lambda kv: kv[1].id):
        for k in sorted(lc.allowed):
            _, c, i, a = k.split("|", 3)
            spec = dom.attr(a)
            tid = lc.items[i].id if i and i in lc.items else lc.id
            for sid, (v, dec) in sorted(lc.allowed[k].items()):
                rows.append([tid, f"{c}.{i}" if i else c, spec.label, sid, spec.show(v), dec])
    _table(ws, 1, ["ID", "対象", "属性", "情報源", "許容する値", "判断"], rows or [["（なし）", "", "", "", "", ""]],
           [10, 40, 12, 16, 16, 12])
    # 未判断のずれ
    ws = wb.create_sheet(_sheet_title("未判断・直し待ち", used))
    rows = [[d.id, f"{d.container}.{d.item}" if d.item else d.container, d.attr.label, SEV_JA[d.severity],
             STATUS_JA[d.status], d.group_label] for d in drifts if d.status != "resolved"]
    _table(ws, 1, ["ずれ", "対象", "属性", "重要度", "状態", "グループ"], rows or [["（なし）", "", "", "", "", ""]],
           [11, 44, 14, 8, 10, 50])
    for w in wb.worksheets:
        if w.title != "表紙":
            w.sheet_view.zoomScale = 90
            if not w.oddFooter.center.text:
                _print_setup(w, f"{dom.title} {w.title}")
    wb.save(path)


def _cell(a, v, d):
    if d and d.status == "open":
        return f"未定（{d.id}）" if v is None else f"{a.show(v)}（{d.id} 未判断）"
    if d and d.status == "decided":
        return f"{a.show(v)}（{d.id} 直し待ち）"
    return a.show(v) if v is not None else "—"


def _review_book(res, versions, out):
    wb = Workbook()
    st = res.state
    n = st.ids["counters"].get("RV", 0) + 1
    st.ids["counters"]["RV"] = n
    st.save()
    rid = f"RV{n:03d}"
    ws = wb.active
    ws.title = "レビュー表"
    docs = list(versions)
    heads = ["指摘No", "対象文書", "版", "箇所（ID）", "箇所の補足（シート・行など）", "指摘内容", "種別", "重要度", "指摘者",
             "対応方針", "対応内容", "反映版", "確認者", "ステータス"]
    ws["A1"], ws["A1"].font = f"レビュー表（{rid}）", _f(True, 14)
    ws["A2"] = "記入するのは「対象文書」〜「指摘者」まで。箇所は各シートのID（TBL-/COL-/TSU-/TC-）で指してください。対応方針以降はbantoが記入します。"
    ws["A2"].font = _f(color="595959")
    rows = []
    vmap = "".join(f'IF(B{{r}}="{d}",{versions[d]},' for d in docs) + '""' + ")" * len(docs)
    for i in range(1, 101):
        rows.append([i, "", f"={vmap.format(r=i + 3)}", "", "", "", "", "", "", "", "", "", "", ""])
    _table(ws, 3, heads, rows, [7, 16, 6, 12, 22, 60, 8, 7, 12, 12, 30, 7, 10, 10])
    ws.freeze_panes = "B4"
    lists = {"B": docs, "G": ["誤り", "漏れ", "表現", "質問", "その他"], "H": ["高", "中", "低"],
             "N": ["未対応", "対応済み", "確認済み"]}
    for col, vals in lists.items():
        dv = DataValidation(type="list", formula1='"' + ",".join(vals) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}4:{col}103")
    # ID一覧（箇所の入力補助）
    wi = wb.create_sheet("ID一覧")
    rows = []
    for dname, ledger in res.ledgers.items():
        dom = get(dname)
        for name, lc in sorted(ledger.items(), key=lambda kv: kv[1].id):
            if lc.status == "removed":
                continue
            rows.append([lc.id, dom.title, name, ""])
            for iname, it in lc.items.items():
                rows.append([it.id, dom.title, name, iname])
    _table(wi, 1, ["ID", "文書", get("tables").container_label + "／" + get("tests").container_label, "項目"], rows,
           [10, 16, 36, 48])
    wi.freeze_panes = "A2"
    dv = DataValidation(type="list", formula1=f"=ID一覧!$A$2:$A${len(rows) + 1}", allow_blank=True)
    dv.error, dv.showErrorMessage = "ID一覧にあるIDを入力してください", True
    ws.add_data_validation(dv)
    dv.add("D4:D103")
    # 前回までの指摘と対応
    wp = wb.create_sheet("これまでの指摘と対応")
    decs = {d.id: d for d in res.decisions}
    rows = []
    for d in res.drifts:
        if d.domain is None:
            r = d.review
            dec = decs.get(d.decision)
            act = {"accept": "対応する", "reject": "対応しない", "answer": "回答"}.get(dec.action, "") if dec else ""
            rows.append([f"{r['doc']}-{r['no']}", r.get("target", ""), r.get("comment", ""), act,
                         dec.reason if dec else "", STATUS_JA[d.status], d.id])
    _table(wp, 1, ["指摘", "箇所", "指摘内容", "対応方針", "対応内容", "状態", "判断票"], rows or [["（なし）"] + [""] * 6],
           [12, 10, 50, 10, 40, 10, 11])
    # bantoが読み戻すための情報
    wm = wb.create_sheet("banto")
    wm["A1"], wm["B1"] = "review_id", rid
    for i, (d, v) in enumerate(versions.items(), 2):
        wm.cell(row=i, column=1, value=f"version:{d}")
        wm.cell(row=i, column=2, value=v)
    wm.sheet_state = "hidden"
    for w in (ws, wi, wp):
        _print_setup(w, f"レビュー表 {rid}")
    path = os.path.join(out, f"review-{rid}.xlsx")
    wb.save(path)
    return path
