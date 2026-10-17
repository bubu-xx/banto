"""記入済みのレビュー表を取り込み、指摘を判断票の流れに乗せる（banto import-review）。"""
import json
import os
import openpyxl
from .state import State
from . import fio

COLS = {"指摘No": "no", "対象文書": "docname", "箇所（ID）": "target", "箇所の補足（シート・行など）": "where",
        "指摘内容": "comment", "種別": "type", "重要度": "severity", "指摘者": "reviewer"}


def import_file(project, path):
    base = os.path.join(project, "banto")
    wb = openpyxl.load_workbook(path, data_only=True)
    if "banto" not in wb.sheetnames or "レビュー表" not in wb.sheetnames:
        raise SystemExit("bantoが出力したレビュー表ではない（「banto」シートがない）")
    meta = wb["banto"]
    rid = str(meta["B1"].value)
    versions = {}
    for row in meta.iter_rows(min_row=2, values_only=True):
        if row[0] and str(row[0]).startswith("version:"):
            versions[str(row[0])[8:]] = row[1]
    ws = wb["レビュー表"]
    header_row = next(r for r in range(1, 10) if ws.cell(row=r, column=1).value == "指摘No")
    heads = {ws.cell(row=header_row, column=c).value: c for c in range(1, ws.max_column + 1)}
    targets = dict(State(base).by_target_id())
    try:
        from . import formats, project as PJ
        targets.update({k: None for k in formats.item_ids(formats.load(base, PJ.load(base).get("format", "standard")))})
    except FileNotFoundError:
        pass
    out, warns = [], []
    for r in range(header_row + 1, ws.max_row + 1):
        rec = {}
        for h, k in COLS.items():
            v = ws.cell(row=r, column=heads[h]).value if h in heads else None
            rec[k] = "" if v is None else str(v).strip()
        if not rec["comment"]:
            continue
        rec["doc"], rec["version"] = rid, versions.get(rec["docname"], "")
        rec["no"] = rec["no"] or str(r - header_row)
        if rec["target"] and rec["target"] not in targets:
            warns.append(f"指摘No.{rec['no']}: 箇所 {rec['target']} は台帳にない")
        if not rec["target"]:
            warns.append(f"指摘No.{rec['no']}: 箇所（ID）が空。判断票には箇所不明として載せる")
        out.append(rec)
    d = os.path.join(base, "reviews")
    os.makedirs(d, exist_ok=True)
    fio.write(os.path.join(d, f"{rid}.json"), json.dumps(out, ensure_ascii=False, indent=1) + "\n")
    return len(out), warns
