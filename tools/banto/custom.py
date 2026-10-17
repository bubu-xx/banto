"""専用様式の取り込み（banto format）。手順は「配置の認識合わせ → 対応づけ → 確定 → 出力」。

  read    様式のExcelを読み、構造（structure.json / STRUCTURE.md）と配置の下書き（layout.json）を作る
  trial   試し刷り。配置の確認用に、各欄へ「［見出し］」を入れたExcel（trial-layout.xlsx）を作る。
          対応づけの後は --data で実際の値を入れたExcel（trial-data.xlsx）を作る
  map     配置の各欄を banto の情報に対応づける下書き（mapping.json）を作る
  confirm 配置（layout）・対応づけ（mapping）を、人が確かめたものとして確定する
配置（どこに何があるか）と対応づけ（そこに何を入れるか）を分けておくので、様式が改版されても配置だけ直せばよい。
下書きは機械的な推定にすぎない。確定するのは人で、AIは確認を手伝う。
"""
import datetime
import json
import os
import re
import shutil
from . import fio
from .extract.doc import HEADERS

FIELD_LABELS = {
    "project.system": ("システム名", "システム", "業務システム名"),
    "project.project": ("案件名", "プロジェクト名", "案件"),
    "project.customer": ("顧客", "顧客名", "お客様", "会社名"),
    "doc.title": ("文書名", "ドキュメント名", "資料名"),
    "doc.version": ("版", "版数", "バージョン", "Ver", "ver", "Version"),
    "doc.date": ("作成日", "日付", "更新日", "最終更新日"),
    "doc.author": ("作成者", "担当者", "作成", "担当"),
    "doc.approver": ("承認者", "承認", "承認日"),
    "container.title": ("テーブル論理名", "テーブル名", "エンティティ名", "テーブル名称", "テストクラス名"),
    "container.name": ("テーブル物理名", "物理テーブル名", "テーブルID", "物理名"),
    "container.summary": ("説明", "概要", "テーブル説明", "目的"),
    "container.id": ("ID", "管理番号"),
}
ITEM_LABELS = {
    "item.no": ("No", "No.", "NO", "番号", "#", "項番"),
    "item.id": ("項目ID", "カラムID"),
    "item.name": HEADERS["name"],
    "item.logical_name": HEADERS["logical_name"],
    "item.base_type": HEADERS["type"],
    "item.length": HEADERS["length"],
    "item.scale": HEADERS["scale"],
    "item.nullable": HEADERS["notnull"] + ("NULL", "NULL可", "NULL許可"),
    "item.pk": HEADERS["pk"] + ("キー", "主キー"),
    "item.fk": ("FK", "外部キー"),
    "item.default": HEADERS["default"],
    "item.indexed": ("INDEX", "インデックス", "IDX"),
    "item.description": HEADERS["description"],
}
HISTORY_LABELS = {
    "history.version": ("版", "版数", "Ver", "バージョン"),
    "history.date": ("日付", "変更日", "更新日", "改訂日"),
    "history.note": ("変更内容", "内容", "改訂内容", "変更箇所", "概要"),
    "history.by": ("変更者", "作成者", "担当", "担当者", "改訂者"),
}


def fdir(base, name):
    return os.path.join(base, "format", name)


def _load(p, default=None):
    return json.loads(fio.read(p)) if os.path.exists(p) else default


def _save(p, obj):
    fio.write(p, json.dumps(obj, ensure_ascii=False, indent=1) + "\n")


def _norm(s):
    return re.sub(r"[\s　:：*※]", "", str(s or ""))


# ---------------------------------------------------------------- 読み取り

def read(base, src, name):
    import openpyxl
    d = fdir(base, name)
    os.makedirs(d, exist_ok=True)
    tpl = os.path.join(d, "template.xlsx")
    if os.path.abspath(src) != os.path.abspath(tpl):
        shutil.copy2(src, tpl)
    wb = openpyxl.load_workbook(tpl)
    structure = {"template": "template.xlsx", "sheets": []}
    layout = {"template": "template.xlsx", "status": "draft", "confirmed_by": "", "sheets": [], "questions": []}
    for ws in wb.worksheets:
        cells = {}
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None or _bordered(c):
                    cells[c.coordinate] = {"v": None if c.value is None else str(c.value), "bold": bool(c.font and c.font.b),
                                           "fill": bool(c.fill and c.fill.fgColor and c.fill.fill_type), "border": _bordered(c),
                                           "row": c.row, "col": c.column}
        merged = [str(m) for m in ws.merged_cells.ranges]
        dvs = [{"cells": str(dv.sqref), "formula": dv.formula1} for dv in ws.data_validations.dataValidation]
        structure["sheets"].append({"sheet": ws.title, "state": ws.sheet_state, "dims": ws.dimensions,
                                    "print_area": str(ws.print_area) if ws.print_area else "", "merged": merged,
                                    "validations": dvs,
                                    "texts": {k: v["v"] for k, v in cells.items() if v["v"] is not None}})
        layout["sheets"].append(_guess(ws, cells, merged, layout["questions"]))
    _save(os.path.join(d, "structure.json"), structure)
    _write_structure_md(d, structure, layout)
    if not os.path.exists(os.path.join(d, "layout.json")) or _load(os.path.join(d, "layout.json"))["status"] != "confirmed":
        _save(os.path.join(d, "layout.json"), layout)
    return layout


def _bordered(c):
    b = c.border
    return bool(b and any(getattr(b, s) and getattr(b, s).style for s in ("left", "right", "top", "bottom")))


def _merged_owner(merged, coord):
    from openpyxl.utils.cell import range_boundaries, coordinate_to_tuple
    r, c = coordinate_to_tuple(coord)
    for m in merged:
        c1, r1, c2, r2 = range_boundaries(m)
        if r1 <= r <= r2 and c1 <= c <= c2:
            return m.split(":")[0]
    return coord


def _guess(ws, cells, merged, questions):
    """配置の下書き：見出しの行（表）と、ラベルと値の組（欄）を推定する。"""
    from openpyxl.utils import get_column_letter
    by_row = {}
    for k, v in cells.items():
        by_row.setdefault(v["row"], []).append(v | {"coord": k})
    tables, used_rows = [], set()
    for r in sorted(by_row):
        texts = [x for x in sorted(by_row[r], key=lambda x: x["col"]) if x["v"]]
        if len(texts) < 3:
            continue
        below = [x for x in by_row.get(r + 1, []) if x["border"]]
        if len(below) < len(texts) - 1 or any(x["v"] for x in by_row.get(r + 1, []) if x["col"] in {t["col"] for t in texts} and x["v"] not in (None, "1")):
            continue
        last = r + 1
        while any(x["border"] for x in by_row.get(last + 1, [])):
            last += 1
        cols = [{"col": get_column_letter(t["col"]), "label": t["v"].strip()} for t in texts]
        tables.append({"id": f"t{len(tables) + 1}", "header_row": r, "first_row": r + 1, "last_row": last,
                       "columns": cols, "confidence": "medium"})
        used_rows |= set(range(r, last + 1))
    fields = []
    for k, v in sorted(cells.items(), key=lambda kv: (kv[1]["row"], kv[1]["col"])):
        if not v["v"] or v["row"] in used_rows:
            continue
        right = get_column_letter(v["col"] + 1) + str(v["row"])
        rc = cells.get(right)
        label = v["v"].strip()
        if len(label) > 20:
            continue
        if rc and not rc["v"] and (rc["border"] or rc["fill"] is False):
            fields.append({"id": f"f{len(fields) + 1}", "label": label, "cell": _merged_owner(merged, right),
                           "confidence": "high" if rc["border"] else "low"})
    title = ws.title
    role = "cover" if "表紙" in title else "history" if re.search("改訂|履歴|変更", title) else \
        ("per-container" if tables and any(re.search("テーブル|エンティティ|物理名", f["label"]) for f in fields) else
         "list" if tables else "other")
    if role == "other" and not fields:
        questions.append(f"シート「{title}」は、欄も表も見つからなかった。説明・注意書きのシートなら other のままでよい")
    for t in tables:
        if t["last_row"] - t["first_row"] < 2:
            questions.append(f"シート「{title}」の表 {t['id']}（{t['header_row']}行目が見出し）は明細の行が少ない。明細が何行まで入るか確認したい")
    return {"sheet": title, "role": role, "fields": fields, "tables": tables}


def _write_structure_md(d, structure, layout):
    L = ["# 様式の読み取り結果", "", "スクリプトが読んだ構造と、配置の下書き（layout.json）。下書きは推定なので、試し刷り（`banto format trial`）で確かめてから確定する。", ""]
    for s, lay in zip(structure["sheets"], layout["sheets"]):
        L += [f"## シート「{s['sheet']}」（推定の役割：{lay['role']}）", "",
              f"- 範囲 {s['dims']}、結合 {len(s['merged'])} か所、入力規則 {len(s['validations'])} か所" +
              (f"、印刷範囲 {s['print_area']}" if s['print_area'] else "") + ("、非表示" if s["state"] != "visible" else ""), ""]
        if lay["fields"]:
            L += ["| 欄 | 見出し | 値を入れるセル | 確度 |", "|---|---|---|---|"]
            L += [f"| {f['id']} | {f['label']} | {f['cell']} | {f['confidence']} |" for f in lay["fields"]]
            L.append("")
        for t in lay["tables"]:
            L += [f"表 {t['id']}：見出し {t['header_row']} 行目、明細 {t['first_row']}〜{t['last_row']} 行目", "",
                  "| 列 | 見出し |", "|---|---|"] + [f"| {c['col']} | {c['label']} |" for c in t["columns"]] + [""]
    if layout["questions"]:
        L += ["## 確認したいこと", ""] + [f"- {q}" for q in layout["questions"]]
    fio.write(os.path.join(d, "STRUCTURE.md"), "\n".join(L) + "\n")


# ---------------------------------------------------------------- 対応づけ

def _match(label, table):
    n = _norm(label)
    for key, names in table.items():
        if n in {_norm(x) for x in names}:
            return key
    return None


def map_draft(base, name, document=None):
    d = fdir(base, name)
    layout = _load(os.path.join(d, "layout.json"))
    if layout is None:
        raise SystemExit("layout.json がない。先に banto format read を実行してください")
    mp = {"status": "draft", "confirmed_by": "",
          "document": document or {"id": "CX", "title": os.path.splitext(name)[0]},
          "fields": {}, "tables": {}, "unmapped": []}
    for s in layout["sheets"]:
        for f in s["fields"]:
            key = _match(f["label"], FIELD_LABELS)
            if s["role"] == "cover" and key and key.startswith("container."):
                key = None
            mp["fields"][f"{s['sheet']}!{f['id']}"] = key or "?"
            if not key:
                mp["unmapped"].append(f"シート「{s['sheet']}」の欄「{f['label']}」（{f['cell']}）")
        for t in s["tables"]:
            source = "history" if s["role"] == "history" else "items" if s["role"] == "per-container" else "containers"
            vocab = HISTORY_LABELS if source == "history" else ITEM_LABELS
            cols = {}
            for c in t["columns"]:
                key = _match(c["label"], vocab)
                if source == "containers" and key in ("item.name", "item.logical_name", "item.description"):
                    key = {"item.name": "container.name", "item.logical_name": "container.title",
                           "item.description": "container.summary"}[key]
                if not key and source == "items":
                    key = f"item.free:{c['label']}"   # bantoに無い情報は、人が書く列として台帳に足す
                    mp["unmapped"].append(f"シート「{s['sheet']}」の表の列「{c['label']}」→ 人が書く列として台帳に追加する案")
                cols[c["col"]] = key or "?"
                if not key:
                    mp["unmapped"].append(f"シート「{s['sheet']}」の表の列「{c['label']}」")
            mp["tables"][f"{s['sheet']}!{t['id']}"] = {"domain": "tables", "source": source, "columns": cols}
    _save(os.path.join(d, "mapping.json"), mp)
    return mp


def confirm(base, name, step, by):
    p = os.path.join(fdir(base, name), f"{step}.json")
    obj = _load(p)
    if obj is None:
        raise SystemExit(f"{step}.json がない")
    if step == "mapping":
        bad = [k for k, v in obj["fields"].items() if v == "?"] + \
              [f"{k}:{c}" for k, t in obj["tables"].items() for c, v in t["columns"].items() if v == "?"]
        if bad:
            raise SystemExit("対応が決まっていない欄がある（? のまま）：" + "、".join(bad) +
                             "。bantoの情報が無い欄は none（空欄のまま）か、表の列なら item.free:<見出し>（人が書く列）にする")
        if _load(os.path.join(fdir(base, name), "layout.json"))["status"] != "confirmed":
            raise SystemExit("先に配置（layout）を確定してください：banto format confirm <名前> --step layout")
    obj["status"], obj["confirmed_by"] = "confirmed", by
    obj["confirmed_at"] = datetime.date.today().isoformat()
    _save(p, obj)
    if step == "mapping":
        _write_format(base, name, obj)
    return obj


def _write_format(base, name, mp):
    """確定した対応づけから、網羅の確認に使う様式定義（format.json）と、台帳に足す「人が書く列」を作る。"""
    from . import project as PJ
    items, extra = [], []
    for key, t in mp["tables"].items():
        cols = [v for v in t["columns"].values() if v.startswith("item.") and v not in ("item.no", "item.id", "item.name")]
        fields = [v[5:] for v in cols]
        extra += [f[5:] for f in fields if f.startswith("free:")]
        if t["source"] == "items":
            cfields = [v[10:] for v in mp["fields"].values() if v.startswith("container.") and v[10:] not in ("name", "id")]
            core = {"logical_name", "base_type", "nullable", "pk", "length", "scale"}
            items.append({"id": f"{mp['document']['id']}-{len(items) + 1:02d}", "title": key.split("!")[0], "kind": "ledger-table",
                          "domain": t["domain"], "fields": fields,
                          "required": [f for f in fields if f in core or f.startswith("free:")],
                          "container_fields": cfields, "container_required": [f for f in cfields if f == "title"],
                          "how": "専用様式の表の列から"})
        elif t["source"] == "containers":
            cf = [v[10:] for v in t["columns"].values() if v.startswith("container.") and v[10:] not in ("name", "id")]
            items.append({"id": f"{mp['document']['id']}-{len(items) + 1:02d}", "title": key.split("!")[0], "kind": "ledger-list",
                          "domain": t["domain"], "fields": cf, "required": cf, "how": "専用様式の一覧の列から"})
    fmt = {"name": f"custom:{name}", "title": f"専用様式「{name}」", "documents": [dict(mp["document"], items=items)]}
    _save(os.path.join(fdir(base, name), "format.json"), fmt)
    cfg = PJ.load(base)
    cur = cfg.setdefault("extra_columns", {}).setdefault("tables", [])
    for c in extra:
        if c not in cur:
            cur.append(c)
    PJ.save(base, cfg)


# ---------------------------------------------------------------- 試し刷りと出力

def trial(base, name, data=False, res=None):
    import openpyxl
    d = fdir(base, name)
    layout = _load(os.path.join(d, "layout.json"))
    wb = openpyxl.load_workbook(os.path.join(d, "template.xlsx"))
    if not data:
        for s in layout["sheets"]:
            ws = wb[s["sheet"]]
            for f in s["fields"]:
                ws[f["cell"]] = f"［{f['label']}］"
            for t in s["tables"]:
                for n in range(3):
                    for c in t["columns"]:
                        ws[f"{c['col']}{t['first_row'] + n}"] = f"［{c['label']}］{n + 1}"
        out = os.path.join(d, "trial-layout.xlsx")
        wb.save(out)
        return out
    mp = _load(os.path.join(d, "mapping.json"))
    out = os.path.join(d, "trial-data.xlsx")
    _fill(res, wb, layout, mp, {}, sample=True)
    wb.save(out)
    return out


def export(res, cfg, out, versions):
    """確定した配置と対応づけで、専用様式に値を入れたExcelを作る。"""
    import openpyxl
    name = cfg["format"].split(":", 1)[1]
    d = fdir(res.base, name)
    layout, mp = _load(os.path.join(d, "layout.json")), _load(os.path.join(d, "mapping.json"))
    if not layout or layout["status"] != "confirmed" or not mp or mp["status"] != "confirmed":
        raise SystemExit(f"専用様式「{name}」の配置・対応づけが確定していない（banto format confirm）")
    wb = openpyxl.load_workbook(os.path.join(d, "template.xlsx"))
    from . import excel
    ver = excel.load_versions(res.base).get(mp["document"]["id"], {})
    _fill(res, wb, layout, mp, ver)
    p = os.path.join(out, f"{mp['document']['id']}_{name}.xlsx")
    wb.save(p)
    return [p]


def _value_of(key, ctx):
    from .coverage import _value
    from .domains import get
    if key in (None, "none", "?"):
        return None
    head, _, rest = key.partition(".")
    if head == "project":
        return ctx["cfg"].get(rest, "")
    if head == "doc":
        return ctx["doc"].get(rest, "")
    if head == "history":
        return ctx["hist"].get(rest, "")
    dom = get(ctx.get("domain", "tables"))
    if head == "container":
        lc = ctx["lc"]
        if rest == "name":
            return lc.name
        if rest == "id":
            return lc.id
        v = lc.attrs.get(rest)
        return "" if v in (None, "") else dom.attrs[rest].show(v)
    if head == "item":
        i = ctx["item"]
        if rest == "no":
            return ctx["no"]
        if rest == "name":
            return i.name
        if rest == "id":
            return i.id
        if rest.startswith("free:"):
            return i.free.get(rest[5:], "")
        v = _value(i, rest)
        if v is None or v == "":
            return ""
        a = dom.attrs[rest]
        if a.vtype == "nullable":                    # 様式の「NOT NULL」欄は ○ で表すことが多い
            return "" if v else "○"
        if a.vtype == "bool":                        # 様式では「該当すれば ○、しなければ空欄」が一般的
            return "○" if v else ""
        if a.vtype == "fk" and not v:
            return ""
        from . import project as PJ
        return PJ.display(ctx["cfg"], a, v)
    return None


def _fill(res, wb, layout, mp, ver, sample=False):
    from copy import copy
    from . import project as PJ
    cfg = PJ.load(res.base)
    hist = ver.get("history", [])
    doc = {"title": mp["document"]["title"], "version": f"第{ver.get('version', 1)}版",
           "date": datetime.date.today().isoformat(), "author": "、".join(PJ.people_with(cfg, "author")),
           "approver": (ver.get("released") or [{}])[-1].get("by", "")}
    conts = sorted((lc for lc in res.ledgers.get("tables", {}).values() if lc.status not in ("removed", "removing")),
                   key=lambda lc: lc.id)
    if sample:
        conts = conts[:1]
    for s in layout["sheets"]:
        ws0 = wb[s["sheet"]]
        targets = [(ws0, None)]
        if s["role"] == "per-container":
            targets = []
            for lc in conts:
                ws = wb.copy_worksheet(ws0)
                ws.title = lc.name[:31]
                targets.append((ws, lc))
            wb.remove(ws0)
        for ws, lc in targets:
            ctx = {"cfg": cfg, "doc": doc, "lc": lc, "domain": "tables"}
            for f in s["fields"]:
                v = _value_of(mp["fields"].get(f"{s['sheet']}!{f['id']}"), ctx)
                if v not in (None, ""):
                    ws[f["cell"]] = v
            for t in s["tables"]:
                tm = mp["tables"].get(f"{s['sheet']}!{t['id']}")
                if not tm:
                    continue
                if tm["source"] == "items" and lc is not None:
                    rows = [{"lc": lc, "item": i, "no": n} for n, i in enumerate(
                        [i for i in lc.items.values() if i.exists is not False], 1)]
                elif tm["source"] == "containers":
                    rows = [{"lc": c, "no": n} for n, c in enumerate(conts, 1)]
                elif tm["source"] == "history":
                    rows = [{"hist": {"version": f"第{h['version']}版", "date": h["date"], "note": h["note"], "by": h.get("by", "")}}
                            for h in hist]
                else:
                    rows = []
                cap = t["last_row"] - t["first_row"] + 1
                for n, r in enumerate(rows):
                    row = t["first_row"] + n
                    for col, key in tm["columns"].items():
                        cell = ws[f"{col}{row}"]
                        if n >= cap:   # 様式の行が足りないときは、明細の1行目の書式を写して続ける
                            src = ws[f"{col}{t['first_row']}"]
                            cell._style = copy(src._style)
                        v = _value_of(key, ctx | r)
                        if v not in (None, ""):
                            cell.value = v
